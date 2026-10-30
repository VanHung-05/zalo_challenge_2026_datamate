import openpyxl
import pandas as pd

from src.datamate.loader import clean_dataframe, load_business_data, parse_vietnamese_date, split_file_stem


def test_parse_date_keeps_time_only_when_asked():
    assert parse_vietnamese_date("2025-04-01 06:00:39") == "2025-04-01"
    assert parse_vietnamese_date("2025-04-01 06:00:39", keep_time=True) == "2025-04-01 06:00:39"
    assert parse_vietnamese_date("01/04/2025 14:05", keep_time=True) == "2025-04-01 14:05:00"
    assert parse_vietnamese_date("2025-04-01 00:00:00", keep_time=True) == "2025-04-01"


def test_parse_date_excel_serial():
    assert parse_vietnamese_date(45717) == "2025-03-01"
    assert parse_vietnamese_date("45748") == "2025-04-01"
    assert parse_vietnamese_date("45717.0") == "2025-03-01"
    # Not a plausible serial: number is kept as is
    assert parse_vietnamese_date(30) == 30


def test_clean_dataframe_drops_repeated_header_rows():
    df = pd.DataFrame({
        "Mã đơn": ["ORD1", "Mã đơn", "ORD2"],
        "Ngày đặt": ["01/01/2025", "Ngày đặt", "02/01/2025"],
        "Tổng tiền": ["10.000", "Tổng tiền", "20.000"],
    })
    out = clean_dataframe(df)
    assert out["Mã đơn"].tolist() == ["ORD1", "ORD2"]
    assert out["Tổng tiền"].tolist() == [10000, 20000]


def test_load_business_merges_same_schema_sheets(tmp_path):
    biz = tmp_path / "biz_x"
    biz.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "T01-2025"
    ws.append(["Mã HĐ", "Thời gian", "Tổng tiền"])
    ws.append(["ORD1", "05/01/2025 08:30", 1000])
    ws2 = wb.create_sheet("T02-2025")
    ws2.append(["Mã HĐ", "Thời gian", "Tổng tiền"])
    ws2.append(["ORD2", "03/02/2025 15:00", 2000])
    wb.save(biz / "HoaDon.xlsx")

    bd = load_business_data(str(biz), use_cache=False, keep_time=True)
    assert list(bd.tables) == ["HoaDon"]
    t = bd.tables["HoaDon"]
    assert t["sheet"].tolist() == ["T01-2025", "T02-2025"]
    assert t["Thời gian"].tolist() == ["2025-01-05 08:30:00", "2025-02-03 15:00:00"]


def test_split_file_stem():
    assert split_file_stem("BanHang_T1_2025") == "BanHang"
    assert split_file_stem("BanHang_Thang01_2026") == "BanHang"
    assert split_file_stem("So_chi_tiet_ban_hang_1") == "So_chi_tiet_ban_hang"
    assert split_file_stem("BanHang_TongHop_2026") == "BanHang_TongHop"
    assert split_file_stem("DonHang_ChiTiet") == "DonHang_ChiTiet"


def _save_sheets(path, sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.save(path)


def test_load_business_merges_table_split_over_files(tmp_path):
    biz = tmp_path / "biz_x"
    biz.mkdir()
    header = ["Mã đơn", "Ngày", "Thành tiền"]
    # A file without a year in its name must still be merged (it used to be left out of BanHang)
    _save_sheets(biz / "So_ban_hang_1.xlsx", {"Sheet1": [header, ["A", "01/05/2025", 10], ["A", "01/05/2025", 10]]})
    _save_sheets(biz / "So_ban_hang_2026.xlsx", {"Sheet1": [header, ["B", "02/01/2026", 20]]})
    _save_sheets(biz / "DanhMuc.xlsx", {"Sheet1": [["Tên", "Giá"], ["X", 1]]})

    bd = load_business_data(str(biz), use_cache=False)
    assert sorted(bd.tables) == ["DanhMuc", "So_ban_hang"]
    # Identical rows are distinct line items: kept
    assert bd.tables["So_ban_hang"]["Mã đơn"].tolist() == ["A", "A", "B"]
    assert bd.metadata["merged_into"] == {"So_ban_hang": ["So_ban_hang_1.xlsx", "So_ban_hang_2026.xlsx"]}


def test_load_business_merges_sheet_split_over_files(tmp_path):
    biz = tmp_path / "biz_x"
    biz.mkdir()
    header = ["Đơn số", "Ngày HĐ", "Tổng thanh toán"]
    _save_sheets(biz / "BanHang_TongHop_2025.xlsx", {
        "BanHang": [header, ["D1", "01/05/2025", 10]],
        "BanHang_HangHoa": [["Mã", "Tên"], ["H1", "Bánh"]],
    })
    _save_sheets(biz / "BanHang_TongHop_2026.xlsx", {"BanHang": [header, ["D2", "02/01/2026", 20]]})

    bd = load_business_data(str(biz), use_cache=False)
    assert sorted(bd.tables) == ["BanHang_TongHop_2025_BanHang_HangHoa", "BanHang_TongHop_BanHang"]
    assert bd.tables["BanHang_TongHop_BanHang"]["Đơn số"].tolist() == ["D1", "D2"]
